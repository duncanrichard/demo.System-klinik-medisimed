from django.shortcuts import render,redirect
from django.http import HttpResponse
import json
from django.core.serializers.json import DjangoJSONEncoder
from urllib.parse import unquote
from inspect import getmembers
from pprint import pprint
from IMMODERMA.globals import Globals
from IMMODERMA.environment import env
from datetime import datetime
from django.conf.urls import url, include
from platform import python_version
import os
from openpyxl import load_workbook
from django.conf import settings
from django.core.files.storage import default_storage
from django.core.files.base import ContentFile


def update_member(request):
    if(Globals().isLogin(request)):
        user_priv = request.session['user_priv']
        user_privelege = request.session['user_priv']
        navbars = Globals().getNavbars(user_priv, 'IMMODERMA', "update_member")
        menubars = Globals().getMenubars(user_priv, 'IMMODERMA', "update_member", '0')
        menubarsChild = Globals().getMenubars(user_priv, 'IMMODERMA', "update_member", '1')
        menubarCount = len(menubars)

        response = render(request, 'dashboard/update_member/base.html', {
            'navbars': navbars,
            'menubars': menubars,
            'menubarsChild': menubarsChild,
            'menubarsType': 1,
            'count_': menubarCount,
            'list_': Globals().getSeparator(menubarCount),
            'user_id': user_privelege
        })
        response['Cache-Control'] = 'no-cache, no-store, max-age=0, must-revalidate'
        return response
    else:
        return redirect('/login')
    
def prosesuploader(request):
	namafile = request.FILES['namafile']
	tmp_file = os.path.join(settings.UPLOAD_PATH, namafile.name)
	if os.path.isfile(tmp_file):
		os.remove(tmp_file)
	path = default_storage.save(tmp_file, ContentFile(namafile.read()))
	file_url = path

	workbook = load_workbook(filename=file_url)
	sheet = workbook.active

	dataJson = {}
	i = 0

	for row in sheet.iter_rows(values_only=True):
		if row[0] not in ['NO',None]:
			data = {
				"NO": row[0],
				"KODE_PASIEN": row[1],
				"MEMBER_ID": row[2],
				"POINT": row[3],
			}
			dataJson[i] = data
			i += 1

	# menampilkan setelah diolah 
	# print(json.dumps(dataJson))

	# menampilkan sesuai kebutuhan
	a = 0
	print (dataJson)
	q = "SET NOCOUNT ON;"
	q += "DECLARE @LIST_GRID FJINKOTAD;"
	for x in dataJson:
		# print(dataJson[a]['no'])
		q += "INSERT INTO @LIST_GRID (FDFJNOM, FDFJPRD_ID,FDFJQTY, FDFJBRG_ID) "
		q += "VALUES ("
		q += str(dataJson[a]['NO'] )+ ","
		q += "'" + dataJson[a]['KODE_PASIEN'] + "',"
		q += str(dataJson[a]['POINT'] )+ ","
		q +=  "'{}');".format(dataJson[a]['MEMBER_ID'])
		a += 1

	q +="select ROW_NUMBER() OVER (ORDER BY FDFJPRD_ID) AS NO,FDFJPRD_ID as KODE_PASIEN,FDFJBRG_ID as MEMBER_ID,FDFJQTY AS POINT,d.NAMAPASIEN "
	q += "from @LIST_GRID b left JOIN PASIEN d "
	q += "on b.FDFJPRD_ID=d.KD_PASIEN "
	# tidak ketemu 
	q +="select ROW_NUMBER() OVER (ORDER BY FDFJPRD_ID) AS NO,FDFJPRD_ID as KODE_PASIEN,FDFJBRG_ID as MEMBER_ID,FDFJQTY AS POINT "
	q += "from @LIST_GRID b "
	q += "where b.FDFJPRD_ID not in (select d.KD_PASIEN from PASIEN d) "
	#  ketemu 
	q +="select ROW_NUMBER() OVER (ORDER BY FDFJPRD_ID) AS NO,FDFJPRD_ID as KODE_PASIEN,FDFJBRG_ID as MEMBER_ID,FDFJQTY AS POINT,d.NAMAPASIEN "
	q += "from @LIST_GRID b  JOIN PASIEN d "
	q += "on b.FDFJPRD_ID=d.KD_PASIEN "
	# cursor.execute(q)
	result1 = Globals().getData(q, None, 1)
	result2 = Globals().getData(q, None, 2)
	result3 = Globals().getData(q, None, 3)

	data = {
		'data1': result1,
		'data2': result2,
		'data3': result3,
	}

	json_data = json.dumps(data, cls=DjangoJSONEncoder)
	return HttpResponse(json_data, content_type="application/json")

def SP_AUD_UPLOAD_MEMBER(request):
    # DETAIL
    NO = json.loads(request.POST['NO'])
    ID_PASIEN = json.loads(request.POST['ID_PASIEN'])
    ID_MEMBER = json.loads(request.POST['ID_MEMBER'])
    POINT = json.loads(request.POST['POINT'])
    # MAIN EXEC
    BUKTI_ID = request.POST['BUKTI_ID']
    TANGGAL = request.POST['TANGGAL']
    KETERANGAN = request.POST['KETERANGAN']
    USERRS = request.POST['USERRS']
    KD_CABANG= request.session['kdCabang']
    StatusAUD = request.POST['StatusAUD']


    q = "SET NOCOUNT ON;DECLARE @LIST_FJINKOTAD FJINKOTAD;"
    q += "DECLARE @NOW datetime; "
    q += "SET @NOW = GETDATE(); "

    proc_param = []
    i = 0
    for x in ID_PASIEN:
        q += "INSERT INTO @LIST_FJINKOTAD (FDFJNOM, FDFJPRD_ID, FDFJBRG_ID, FDFJQTY) VALUES ("
        q += "'" + NO[i] + "',"
        q += "'" + ID_PASIEN[i] + "',"
        q += "'" + ID_MEMBER[i] + "',"
        q += POINT[i] + ");"
        i += 1

    q += "EXEC AUD_UPLOAD_MEMBER "
    q += "'" + BUKTI_ID + "',"
    q += "'" + TANGGAL + "',"
    q += "'" + KETERANGAN + "',"
    q += "'" + USERRS + "',"
    q += "@NOW,"
    q += "'" + KD_CABANG + "',"
    q += "'" + StatusAUD + "',"
    q += "@LIST_FJINKOTAD"
    if (StatusAUD=='D') :
        user = {
        'user_id': request.session['user_id'],
        'user_name': request.session['user_name'],
        'user_priv': request.session['user_priv'],
        }
        data = []
        data.append({"query": "select * from UPDATE_MEMBERH WHERE FHUMBUKTI_ID = '" + BUKTI_ID + "'"})
        data.append({"query": "select * from UPDATE_MEMBERD WHERE FDUMBUKTI_ID = '" + BUKTI_ID + "'"})
        Globals().create_log('Hapus LogDelete.txt', 'IMMODERMA', data, user)
  
    try:
        result = Globals().getDataSP(q)
        json_data = json.dumps(result, cls=DjangoJSONEncoder)
        return HttpResponse(json_data, content_type="application/json")

    except ValueError:
        print(q)

def getUploadMember(request):
	no_bukti = '%'+request.GET['no_bukti']+'%'
	keterangan = '%'+request.GET['keterangan']+'%'
	tipe = request.GET['tipe']
	tanggal = datetime.strptime(request.GET['tanggal'], "%Y-%m-%d")

	if(tipe == 'mutasi_by_bulan'):
		q = "select top 100 FHUMBUKTI_ID, FHUMKETERANGAN, convert(varchar, FHUMTANGGAL, 23) as TANGGAL "
		q +="from UPDATE_MEMBERH where (FHUMBUKTI_ID like %s) and "
		q += "(FHUMKETERANGAN like %s) and (YEAR(FHUMTANGGAL) = %s) and (MONTH(FHUMTANGGAL) = %s) "
		result = Globals().getDataQuery(q, [no_bukti, keterangan, tanggal.year, tanggal.month])
	else:
		q = "select top 100 FHUMBUKTI_ID, FHUMKETERANGAN, convert(varchar, FHUMTANGGAL, 23) as TANGGAL "
		q +="from UPDATE_MEMBERH where (FHUMBUKTI_ID like %s)  and "
		q += "(FHUMKETERANGAN like %s) and (FHUMTANGGAL = %s)"
		result = Globals().getDataQuery(q, [no_bukti, keterangan, tanggal])

	json_data = json.dumps(result, cls=DjangoJSONEncoder)
	return HttpResponse(json_data, content_type="application/json")

def getUploadMemberByBukti(request):
    no_bukti = request.GET['no_bukti']
    q = "select FHUMBUKTI_ID, FHUMKETERANGAN, convert(varchar, FHUMTANGGAL, 23) as TANGGAL, "
    q += "B.FDUMPASIEN_ID as KODE_PASIEN, FDUMMEMBER_ID as MEMBER_ID, FDUMPOINT as POINT, FDUMNO as NO,C.NAMAPASIEN,ISNULL(D.KETERANGAN,'') AS KETERANGAN "
    q += "from UPDATE_MEMBERH A  "
    q += "INNER JOIN UPDATE_MEMBERD B ON A.FHUMBUKTI_ID=B.FDUMBUKTI_ID "
    q += "LEFT JOIN PASIEN C ON B.FDUMPASIEN_ID=C.KD_PASIEN "
    q += "LEFT JOIN MEMBER D ON B.FDUMMEMBER_ID=D.KODE_MEMBER "
    q += "where a.FHUMBUKTI_ID = %s order by FDUMNO asc "
    # tidak ketemu 
    q += "select FHUMBUKTI_ID, FHUMKETERANGAN, convert(varchar, FHUMTANGGAL, 23) as TANGGAL, "
    q += "B.FDUMPASIEN_ID as KODE_PASIEN, FDUMMEMBER_ID as MEMBER_ID, FDUMPOINT as POINT, FDUMNO as NO "
    q += "from UPDATE_MEMBERH A  "
    q += "INNER JOIN UPDATE_MEMBERD B ON A.FHUMBUKTI_ID=B.FDUMBUKTI_ID "
    q += "where a.FHUMBUKTI_ID = %s AND b.FDUMPASIEN_ID not in (select d.KD_PASIEN from PASIEN d)  order by FDUMNO asc "
    #  ketemu 
    q += "select FHUMBUKTI_ID, FHUMKETERANGAN, convert(varchar, FHUMTANGGAL, 23) as TANGGAL, "
    q += "B.FDUMPASIEN_ID as KODE_PASIEN, FDUMMEMBER_ID as MEMBER_ID, FDUMPOINT as POINT, FDUMNO as NO,C.NAMAPASIEN,ISNULL(D.KETERANGAN,'') AS KETERANGAN "
    q += "from UPDATE_MEMBERH A  "
    q += "INNER JOIN UPDATE_MEMBERD B ON A.FHUMBUKTI_ID=B.FDUMBUKTI_ID "
    q += "INNER JOIN PASIEN C ON B.FDUMPASIEN_ID=C.KD_PASIEN "
    q += "LEFT JOIN MEMBER D ON B.FDUMMEMBER_ID=D.KODE_MEMBER "
    q += "where a.FHUMBUKTI_ID = %s order by FDUMNO asc"
    result1 = Globals().getData(q, [no_bukti,no_bukti,no_bukti], 1)
    result2 = Globals().getData(q, [no_bukti,no_bukti,no_bukti], 2)
    result3 = Globals().getData(q, [no_bukti,no_bukti,no_bukti], 3)

    data = {
        'data1': result1,
        'data2': result2,
        'data3': result3,
    }

    json_data = json.dumps(data, cls=DjangoJSONEncoder)
    return HttpResponse(json_data, content_type="application/json")

