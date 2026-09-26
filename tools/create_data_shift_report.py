from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\DUNCAN\Desktop\Project_Free\IMMODERMA")
OUTPUT = ROOT / "Laporan_Analisis_Pergeseran_Data_Member_IMMODERMA.docx"

NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
RED = RGBColor(156, 0, 6)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    return doc.add_paragraph(text, style=style)


def add_number(doc, text):
    return doc.add_paragraph(text, style="List Number")


def add_manual_number(doc, number, text):
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.add_run(f"{number}. ").bold = True
    p.add_run(text)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    for i, line in enumerate(text.splitlines()):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(8.5)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(8.5)
        if widths:
            cell.width = widths[i]
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(8.3)
            if i == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if widths:
                cell.width = widths[i]
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.1)
section.bottom_margin = Cm(1.9)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = BLACK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

body = styles["Body Text"]
body.font.name = "Aptos"
body._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
body._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
body.font.size = Pt(10.5)
body.font.color.rgb = BLACK
body.paragraph_format.space_after = Pt(7)
body.paragraph_format.line_spacing = 1.12

title = styles["Title"]
title.font.name = "Aptos Display"
title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title.font.size = Pt(25)
title.font.bold = True
title.font.color.rgb = BLACK
title_ppr = title._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = BLACK
    st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
    st.paragraph_format.space_after = Pt(5)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("Laporan Analisis Pergeseran Data Member IMMODERMA")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Analisis awal penyebab ketidaksesuaian nomor member dan identitas pasien")
r.font.size = Pt(13)
r.font.color.rgb = RGBColor(70, 70, 70)

meta = add_table(
    doc,
    ["Keterangan", "Isi"],
    [
        ("Sistem", "SIMKLINIK IMMODERMA"),
        ("Objek analisis", "Modul update member dan alur upload Excel"),
        ("Tanggal laporan", "26 September 2026"),
        ("Status", "Analisis teknis awal untuk klarifikasi vendor"),
    ],
    [Cm(4.2), Cm(11.4)],
)

add_body(
    doc,
    "Laporan ini menilai kemungkinan penyebab data nomor member tampak atau benar-benar berpindah antar pasien. Fokus analisis adalah contoh pola ketika member 100 seharusnya terkait dengan Duncan dan member 123 seharusnya terkait dengan Yusuf, tetapi sistem kemudian menampilkan Duncan bersama member 123.",
)
add_body(
    doc,
    "Kesimpulan utama: bukti kode yang tersedia belum cukup untuk menyatakan bahwa ketidaksesuaian tersebut terjadi karena tindakan manual user. Aplikasi memiliki beberapa jalur teknis yang dapat menghasilkan atau mempertahankan pasangan pasien dan nomor member yang salah, sementara pencatatan audit untuk proses tambah atau pembaruan member tidak terlihat pada kode aplikasi yang diperiksa.",
    bold_lead="Kesimpulan utama:",
)

doc.add_page_break()

add_heading(doc, "Ringkasan Eksekutif", 1)
add_body(
    doc,
    "Modul update member menerima data dalam bentuk pasangan KODE_PASIEN dan MEMBER_ID dari Excel. Nama pasien tidak menjadi sumber validasi kepemilikan member; nama hanya diambil kemudian dari tabel PASIEN berdasarkan KODE_PASIEN. Oleh karena itu, jika satu baris input berisi kode pasien Duncan dan member 123, aplikasi akan menampilkan Duncan dan member 123 tanpa membuktikan bahwa member 123 benar-benar milik Duncan.",
)
add_body(
    doc,
    "Pada tahap simpan, browser mengirim empat array paralel yaitu nomor, kode pasien, member ID, dan poin. Server memasukkan nilai tersebut ke table variable lalu menjalankan stored procedure AUD_UPLOAD_MEMBER. Definisi stored procedure tidak berada dalam repository, sehingga bagian paling menentukan dalam pembaruan database belum dapat diaudit dari kode sumber ini.",
)
add_body(
    doc,
    "Kode aplikasi mencatat file log untuk operasi hapus ketika StatusAUD bernilai D. Tidak ditemukan pencatatan sebelum dan sesudah untuk operasi tambah atau pembaruan member. Nilai user session memang diteruskan ke stored procedure sebagai USERRS, tetapi nilai itu hanya menunjukkan sesi yang memulai permintaan apabila benar-benar disimpan oleh stored procedure; nilai tersebut tidak dengan sendirinya membuktikan bahwa user sengaja menukar data.",
)

add_heading(doc, "Penilaian Kesimpulan Vendor", 2)
add_table(
    doc,
    ["Pernyataan", "Penilaian", "Alasan"],
    [
        ("Data berubah karena di-update user", "Belum terbukti", "Tidak tersedia audit before-after untuk tambah atau update pada kode aplikasi yang diperiksa."),
        ("Aplikasi tidak mungkin menyebabkan pergeseran", "Tidak dapat diterima", "Kode mempercayai pasangan kode pasien dan member dari input serta meneruskannya ke stored procedure tanpa validasi kepemilikan."),
        ("User session tercatat", "Mungkin", "USERRS diteruskan ke stored procedure, tetapi perlu dibuktikan bahwa nilainya disimpan bersama nilai lama, nilai baru, waktu, dan sumber transaksi."),
    ],
    [Cm(4.1), Cm(3.0), Cm(8.5)],
)

add_heading(doc, "Ruang Lingkup dan Sumber Bukti", 1)
add_bullet(doc, "Kode Python modul dashboard/update_member.py.")
add_bullet(doc, "Kode JavaScript templates/dashboard/update_member/js/init.html dan upluadexcell.html.")
add_bullet(doc, "Fungsi logging IMMODERMA/globals.py.")
add_bullet(doc, "Struktur pemanggilan stored procedure AUD_UPLOAD_MEMBER dari aplikasi.")
add_bullet(doc, "Pola tabel UPDATE_MEMBERH, UPDATE_MEMBERD, PASIEN, dan MEMBER yang terlihat dari query aplikasi.")
add_body(
    doc,
    "Analisis ini belum mencakup isi stored procedure AUD_UPLOAD_MEMBER, constraint database, trigger, SQL Server Audit, Change Data Capture, transaction log, maupun isi file Excel pada saat kejadian. Karena itu, laporan membedakan temuan terverifikasi dari hipotesis yang masih perlu diuji.",
)

add_heading(doc, "Alur Data yang Terverifikasi", 1)
add_number(doc, "Aplikasi membaca worksheet aktif dari file Excel.")
add_number(doc, "Setiap baris dipetakan: kolom pertama sebagai nomor, kolom kedua sebagai kode pasien, kolom ketiga sebagai member ID, dan kolom keempat sebagai poin.")
add_number(doc, "Preview mencari nama pasien dengan join berdasarkan kode pasien saja.")
add_number(doc, "Preview membuat ulang nomor urut dengan ROW_NUMBER berdasarkan kode pasien.")
add_number(doc, "Browser mengambil baris yang terlihat dan membuat array NO, ID_PASIEN, ID_MEMBER, serta POINT.")
add_number(doc, "Server memasukkan array tersebut berdasarkan indeks yang sama ke table variable FJINKOTAD.")
add_number(doc, "Server menjalankan stored procedure AUD_UPLOAD_MEMBER untuk menulis hasil ke database.")

add_code(
    doc,
    "Excel:     [NO] [KODE_PASIEN] [MEMBER_ID] [POINT]\n"
    "Preview:   KODE_PASIEN -> JOIN PASIEN -> NAMAPASIEN\n"
    "Simpan:    arrays paralel -> table variable -> AUD_UPLOAD_MEMBER\n"
    "Baca ulang: UPDATE_MEMBERD -> JOIN PASIEN dan MEMBER",
)

add_heading(doc, "Temuan Teknis dan Indikasi Pergeseran", 1)
findings = [
    (
        "F1",
        "Tidak ada validasi kepemilikan member",
        "Aplikasi menerima MEMBER_ID pada baris yang sama dengan KODE_PASIEN dan hanya memeriksa bahwa keduanya tidak kosong. Tidak ada pengecekan bahwa member tersebut sebelumnya atau secara sah dimiliki pasien tersebut.",
        "Tinggi",
    ),
    (
        "F2",
        "Nomor baris input diganti pada preview",
        "Query preview menggunakan ROW_NUMBER OVER ORDER BY FDFJPRD_ID. Urutan asli Excel tidak dipertahankan sebagai identitas audit sehingga data terlihat berpindah baris ketika dibandingkan berdasarkan nomor urut.",
        "Sedang",
    ),
    (
        "F3",
        "Logika tulis utama berada di stored procedure",
        "AUD_UPLOAD_MEMBER tidak tersedia dalam repository. Kesalahan join, update tanpa key yang unik, urutan cursor, cakupan cabang, atau transaksi di procedure belum dapat dikesampingkan.",
        "Tinggi",
    ),
    (
        "F4",
        "Audit tambah dan update tidak terlihat",
        "Kode create_log dipanggil hanya ketika StatusAUD sama dengan D untuk operasi hapus. Tidak terlihat snapshot nilai lama dan baru untuk tambah atau pembaruan member.",
        "Tinggi",
    ),
    (
        "F5",
        "File upload dengan nama sama dapat ditimpa",
        "Sebelum menyimpan upload, aplikasi menghapus file lama jika path dengan nama yang sama sudah ada. Bukti sumber kejadian dapat hilang jika nama file digunakan kembali.",
        "Sedang",
    ),
    (
        "F6",
        "Constraint database belum terverifikasi",
        "Belum ada bukti bahwa MEMBER_ID unik per pasien atau bahwa pasangan pasien-member dilindungi foreign key dan unique constraint.",
        "Tinggi jika constraint tidak ada",
    ),
    (
        "F7",
        "Duplikasi kode pasien atau member belum dikesampingkan",
        "Join berdasarkan KD_PASIEN dan KODE_MEMBER dapat menghasilkan hasil ambigu apabila key tidak unik atau data master sudah duplikat.",
        "Sedang",
    ),
    (
        "F8",
        "Input Excel dapat sudah tidak sejajar",
        "Jika satu kolom Excel diurutkan, ditempel, atau dihapus terpisah dari kolom lain, aplikasi tetap menganggap setiap baris sebagai pasangan yang sah.",
        "Tinggi",
    ),
]
add_table(
    doc,
    ["ID", "Indikasi", "Penjelasan", "Risiko"],
    findings,
    [Cm(1.0), Cm(4.0), Cm(8.2), Cm(2.4)],
)

add_heading(doc, "Cara Membedakan Jenis Pergeseran", 1)
add_table(
    doc,
    ["Gejala", "Lokasi indikasi", "Interpretasi awal"],
    [
        ("Preview upload sudah menunjukkan Duncan dengan member 123", "Excel atau parser upload", "Pasangan salah sudah ada sebelum stored procedure dijalankan."),
        ("Preview benar tetapi setelah simpan atau dibuka ulang menjadi salah", "Stored procedure atau database", "Perubahan terjadi pada tahap tulis, trigger, atau pembacaan ulang."),
        ("Pasangan pasien-member tetap benar tetapi nomor baris berubah", "Query preview", "Ini efek pengurutan ulang ROW_NUMBER, bukan bukti pertukaran identitas."),
        ("Satu member muncul pada beberapa pasien", "Integritas database", "Kemungkinan tidak ada unique constraint, ada duplikasi, atau procedure melakukan update tidak spesifik."),
        ("Kejadian hanya muncul pada cabang atau batch tertentu", "Cakupan procedure", "Periksa penggunaan KD_CABANG, nomor bukti, dan filter update di stored procedure."),
        ("Data berubah tanpa catatan nilai lama", "Audit aplikasi", "Tidak cukup bukti untuk menentukan pelaku atau mekanisme perubahan."),
    ],
    [Cm(5.0), Cm(3.6), Cm(7.0)],
)

add_heading(doc, "Mengapa Klaim Update oleh User Belum Cukup", 1)
add_body(
    doc,
    "Keberadaan user ID pada transaksi tidak otomatis membuktikan bahwa user secara manual menukar nomor member. User ID dapat menunjukkan siapa yang menekan simpan atau menjalankan upload, sedangkan pasangan yang salah dapat berasal dari file Excel, pengurutan preview, logika stored procedure, trigger database, atau data master yang sudah tidak konsisten.",
)
add_body(
    doc,
    "Atribusi kepada user baru layak dibuat jika tersedia bukti yang menghubungkan satu identitas user dengan waktu kejadian, nilai sebelum, nilai sesudah, kode pasien, member ID, sumber transaksi, dan operasi SQL yang dilakukan. Tanpa rangkaian tersebut, kesimpulan mengenai pelaku bersifat asumsi.",
)

add_heading(doc, "Bukti Minimum untuk Atribusi", 2)
add_bullet(doc, "Identitas user aplikasi dan identitas login SQL Server.")
add_bullet(doc, "Timestamp presisi saat nilai berubah.")
add_bullet(doc, "Nilai lama dan nilai baru untuk KODE_PASIEN serta MEMBER_ID.")
add_bullet(doc, "Nomor bukti atau batch upload yang terkait.")
add_bullet(doc, "File Excel asli beserta hash atau salinan yang tidak dapat ditimpa.")
add_bullet(doc, "Definisi dan jejak eksekusi stored procedure AUD_UPLOAD_MEMBER.")
add_bullet(doc, "SQL Server Audit, Extended Events, CDC, temporal table, trigger audit, atau transaction log yang relevan.")

doc.add_page_break()

add_heading(doc, "Pemeriksaan Lanjutan yang Disarankan", 1)
add_heading(doc, "Pemeriksaan Stored Procedure", 2)
add_body(doc, "Ambil definisi procedure dan periksa seluruh perintah INSERT, UPDATE, MERGE, cursor, join, filter, serta penggunaan nomor urut dan KD_CABANG.")
add_code(doc, "EXEC sp_helptext 'dbo.AUD_UPLOAD_MEMBER';")

add_heading(doc, "Pemeriksaan Duplikasi Member", 2)
add_code(
    doc,
    "SELECT FDUMMEMBER_ID,\n"
    "       COUNT(DISTINCT FDUMPASIEN_ID) AS JUMLAH_PASIEN\n"
    "FROM UPDATE_MEMBERD\n"
    "GROUP BY FDUMMEMBER_ID\n"
    "HAVING COUNT(DISTINCT FDUMPASIEN_ID) > 1;",
)

add_heading(doc, "Pemeriksaan Duplikasi Pasien", 2)
add_code(
    doc,
    "SELECT KD_PASIEN, COUNT(*) AS JUMLAH\n"
    "FROM PASIEN\n"
    "GROUP BY KD_PASIEN\n"
    "HAVING COUNT(*) > 1;",
)

add_heading(doc, "Pemeriksaan Pasangan per Nomor Bukti", 2)
add_code(
    doc,
    "SELECT D.FDUMBUKTI_ID, D.FDUMNO, D.FDUMPASIEN_ID,\n"
    "       P.NAMAPASIEN, D.FDUMMEMBER_ID\n"
    "FROM UPDATE_MEMBERD D\n"
    "LEFT JOIN PASIEN P ON D.FDUMPASIEN_ID = P.KD_PASIEN\n"
    "ORDER BY D.FDUMBUKTI_ID, D.FDUMNO;",
)

add_heading(doc, "Pemeriksaan Constraint", 2)
add_code(
    doc,
    "SELECT fk.name AS foreign_key_name,\n"
    "       OBJECT_NAME(fk.parent_object_id) AS parent_table,\n"
    "       OBJECT_NAME(fk.referenced_object_id) AS referenced_table\n"
    "FROM sys.foreign_keys fk\n"
    "WHERE OBJECT_NAME(fk.parent_object_id) IN ('UPDATE_MEMBERD','PASIEN','MEMBER');",
)

add_heading(doc, "Uji Reproduksi Terkendali", 2)
add_manual_number(doc, 1, "Buat dua pasien uji dan dua nomor member uji yang belum pernah digunakan.")
add_manual_number(doc, 2, "Simpan salinan Excel asli dan hitung hash file sebelum upload.")
add_manual_number(doc, 3, "Rekam pasangan pada Excel, preview, payload browser, table variable, hasil stored procedure, dan hasil baca ulang.")
add_manual_number(doc, 4, "Ulangi dengan urutan Excel dibalik untuk melihat apakah hasil bergantung pada posisi baris.")
add_manual_number(doc, 5, "Ulangi pada dua cabang jika sistem multi-cabang untuk menguji isolasi KD_CABANG.")
add_manual_number(doc, 6, "Bandingkan seluruh tahap untuk menemukan titik pertama pasangan berubah.")

doc.add_page_break()
add_heading(doc, "Rekomendasi Pengendalian", 1)
add_table(
    doc,
    ["Prioritas", "Rekomendasi", "Tujuan"],
    [
        ("P1", "Tambahkan validasi bahwa satu MEMBER_ID hanya boleh terkait dengan pasien yang sah sebelum stored procedure dipanggil.", "Mencegah pasangan salah masuk ke proses tulis."),
        ("P1", "Tambahkan audit before-after untuk setiap tambah, update, dan hapus, termasuk user, timestamp, nomor bukti, cabang, pasien, serta member.", "Membuktikan mekanisme dan pelaku perubahan."),
        ("P1", "Audit dan uji AUD_UPLOAD_MEMBER; pastikan seluruh update memakai key pasien dan member yang eksplisit serta transaction yang konsisten.", "Menghilangkan risiko pergeseran akibat urutan atau join."),
        ("P2", "Pertahankan nomor baris asli Excel dan tambahkan ID baris teknis yang tidak berubah selama preview dan simpan.", "Menjaga traceability antar tahap."),
        ("P2", "Simpan file upload secara immutable dengan nama unik dan hash file.", "Menjaga bukti input asli."),
        ("P2", "Tambahkan unique constraint atau aturan integritas yang sesuai setelah data duplikat dibersihkan.", "Mencegah satu member terhubung ke beberapa pasien tanpa aturan bisnis yang sah."),
        ("P3", "Gunakan parameterized query untuk menggantikan perakitan SQL berbasis string.", "Meningkatkan keamanan dan keandalan pemrosesan nilai input."),
    ],
    [Cm(1.7), Cm(9.1), Cm(4.8)],
)

add_heading(doc, "Kesimpulan", 1)
add_body(
    doc,
    "Berdasarkan kode yang diperiksa, terdapat mekanisme teknis yang memungkinkan pasangan pasien dan nomor member salah diterima, ditampilkan, atau disimpan tanpa validasi kepemilikan. Terdapat pula perubahan nomor urut pada preview dan ketergantungan pada stored procedure yang belum diperiksa. Dengan kondisi audit saat ini, tidak ada dasar yang cukup untuk menyimpulkan bahwa pergeseran pasti disebabkan oleh pembaruan manual user.",
)
add_body(
    doc,
    "Vendor perlu menunjukkan bukti audit yang lengkap apabila tetap menyatakan perubahan dilakukan oleh user. Pada saat yang sama, tim teknis perlu memperoleh definisi AUD_UPLOAD_MEMBER dan memeriksa constraint database untuk menentukan titik perubahan secara pasti. Sampai bukti tersebut tersedia, status yang tepat adalah penyebab belum terkonfirmasi dengan beberapa indikasi teknis pada jalur upload dan penyimpanan.",
)

doc.add_page_break()
add_heading(doc, "Referensi Kode", 1)
references = [
    ("R1", "dashboard/update_member.py baris 57 sampai 65", "Pembacaan kolom Excel menjadi NO, KODE_PASIEN, MEMBER_ID, dan POINT."),
    ("R2", "dashboard/update_member.py baris 86 sampai 98", "Join nama pasien berdasarkan KODE_PASIEN dan pembuatan ulang nomor urut."),
    ("R3", "dashboard/update_member.py baris 111 sampai 148", "Pembuatan table variable dan pemanggilan AUD_UPLOAD_MEMBER."),
    ("R4", "dashboard/update_member.py baris 149 sampai 158", "Logging hanya pada StatusAUD D atau operasi hapus."),
    ("R5", "dashboard/update_member.py baris 188 sampai 213", "Pembacaan ulang pasangan dari UPDATE_MEMBERD serta join PASIEN dan MEMBER."),
    ("R6", "templates/dashboard/update_member/js/init.html baris 49 sampai 59", "Pembentukan array paralel dari visible rows untuk proses simpan."),
    ("R7", "IMMODERMA/globals.py baris 221 sampai 260", "Implementasi file log dan snapshot query yang diberikan pemanggil."),
]
add_table(doc, ["ID", "Lokasi", "Relevansi"], references, [Cm(1.0), Cm(6.2), Cm(8.4)])

# Footer
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Laporan Analisis Pergeseran Data Member IMMODERMA")
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)

# Metadata
doc.core_properties.title = "Laporan Analisis Pergeseran Data Member IMMODERMA"
doc.core_properties.subject = "Analisis teknis ketidaksesuaian nomor member dan identitas pasien"
doc.core_properties.author = "Analisis teknis IMMODERMA"

doc.save(OUTPUT)
print(OUTPUT)
